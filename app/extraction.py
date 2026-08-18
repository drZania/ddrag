"""Pure document text-extraction helpers for M5 step 1."""

from __future__ import annotations

import io

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "text/plain",
    "text/markdown",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExtractionError(Exception):
    """Raised when text cannot be extracted from the provided document bytes."""


def _extract_plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ExtractionError("Document is not valid UTF-8 text") from exc


def _extract_pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError) as exc:
        raise ExtractionError("Failed to read PDF document") from exc
    return "\n".join(pages_text)


def _extract_docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Failed to read DOCX document") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(content_type: str, data: bytes) -> str:
    """Extract plain text from document bytes based on the provided MIME type."""

    if content_type in ("text/plain", "text/markdown"):
        return _extract_plain_text(data)
    if content_type == "application/pdf":
        return _extract_pdf_text(data)
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx_text(data)

    raise ExtractionError(f"Unsupported content type for extraction: {content_type}")
