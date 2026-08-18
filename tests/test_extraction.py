"""Focused tests for the pure M5 step 1 text-extraction module."""

import io

import docx
import pytest

from app.extraction import ExtractionError, extract_text


def _build_pdf_bytes(text: str) -> bytes:
    # pypdf's writer cannot draw text without extra dependencies, so build a
    # minimal valid PDF by hand, including a proper xref table and trailer,
    # whose content stream pypdf can extract text from.
    content = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    header = b"%PDF-1.4\n"
    body = bytearray(header)
    offsets = [0]
    for index, obj_body in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{index} 0 obj\n".encode("latin-1") + obj_body + b"\nendobj\n"

    xref_offset = len(body)
    body += f"xref\n0 {len(objects) + 1}\n".encode("latin-1")
    body += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        body += f"{offset:010} 00000 n \n".encode("latin-1")
    body += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode("latin-1")

    return bytes(body)


def _build_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row_values in enumerate(table_rows):
            for col_index, value in enumerate(row_values):
                table.cell(row_index, col_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_plain_text() -> None:
    result = extract_text("text/plain", b"hello world")

    assert result == "hello world"


def test_extract_markdown_text() -> None:
    result = extract_text("text/markdown", b"# Title\n\nSome body text.")

    assert result == "# Title\n\nSome body text."


def test_extract_plain_text_unicode() -> None:
    payload = "cafe\u0301 \u00e9\u00e8\u00ea \u4f60\u597d".encode("utf-8")

    result = extract_text("text/plain", payload)

    assert result == payload.decode("utf-8")


def test_extract_plain_text_rejects_invalid_utf8() -> None:
    invalid_bytes = b"\xff\xfe not valid utf-8"

    with pytest.raises(ExtractionError):
        extract_text("text/plain", invalid_bytes)


def test_extract_plain_text_empty_document() -> None:
    result = extract_text("text/plain", b"")

    assert result == ""


def test_extract_pdf_text() -> None:
    pdf_bytes = _build_pdf_bytes("Hello PDF")

    result = extract_text("application/pdf", pdf_bytes)

    assert "Hello PDF" in result


def test_extract_pdf_rejects_malformed_bytes() -> None:
    with pytest.raises(ExtractionError):
        extract_text("application/pdf", b"not a real pdf")


def test_extract_docx_paragraph_text() -> None:
    docx_bytes = _build_docx_bytes(["First paragraph.", "Second paragraph."])

    result = extract_text(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_bytes,
    )

    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_extract_docx_table_text() -> None:
    docx_bytes = _build_docx_bytes(
        ["Intro paragraph."],
        table_rows=[["Header A", "Header B"], ["Row1 A", "Row1 B"]],
    )

    result = extract_text(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_bytes,
    )

    assert "Header A" in result
    assert "Row1 B" in result


def test_extract_docx_rejects_invalid_bytes() -> None:
    with pytest.raises(ExtractionError):
        extract_text(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not a real docx",
        )


def test_extract_text_rejects_unsupported_content_type() -> None:
    with pytest.raises(ExtractionError):
        extract_text("application/octet-stream", b"bytes")
